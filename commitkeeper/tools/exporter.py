"""Tool 4 - persistent review views: Excel register and an .ics calendar.

SQLite stays authoritative. The workbook is regenerated from the store on every
export, except the two manual columns ("Your decision", "Your comments"): those
are read back first, stored in `manual_reviews`, and written into the new file,
so an export can never overwrite what the user typed.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..logbook import utc_now
from ..memory import Memory

REGISTER_COLUMNS = [
    "Commitment ID", "Direction", "Owner", "Counterparty", "What", "Current due date", "Due phrase",
    "Superseded due dates", "Status", "Status reason", "Agent recommended next action", "Decided by",
    "Recommendation reason", "Evidence (message: quote)", "Nudges sent", "Last contact outcome",
    "Pending approval", "Your decision", "Your comments", "Updated (UTC)",
]
MANUAL_COLUMNS = ("Your decision", "Your comments")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
MANUAL_FILL = PatternFill("solid", fgColor="FFF2CC")


def _read_manual(path: Path) -> dict[str, tuple[str, str]]:
    if not path.exists():
        return {}
    ws = load_workbook(path)["Register"]
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) for name in ("Commitment ID", *MANUAL_COLUMNS) if name in header}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = row[idx["Commitment ID"]] if "Commitment ID" in idx else None
        if cid:
            out[cid] = tuple(str(row[idx[c]] or "").strip() if c in idx else "" for c in MANUAL_COLUMNS)
    return out


def _sync_manual(mem: Memory, from_file: dict[str, tuple[str, str]]) -> dict[str, tuple[str, str]]:
    """Store changed manual entries (history kept) and return the latest per commitment."""
    conn = mem.conn
    latest = {}
    for row in conn.execute("SELECT * FROM manual_reviews WHERE profile_id = ? ORDER BY review_id",
                            (mem.profile_id,)):
        latest[row["commitment_id"]] = (row["user_decision"] or "", row["user_comments"] or "")
    for cid, values in from_file.items():
        if any(values) and latest.get(cid) != values:
            conn.execute("INSERT INTO manual_reviews (profile_id, commitment_id, user_decision, user_comments,"
                         " imported_at) VALUES (?, ?, ?, ?, ?)", (mem.profile_id, cid, *values, utc_now()))
            mem.log.event("memory_write", "manual_reviews.import", {"commitment_id": cid, "values": values})
            latest[cid] = values
    conn.commit()
    return latest


def _style(ws, widths: dict[int, int], manual_cols: tuple[int, ...] = ()) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col, width in widths.items():
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column in manual_cols:
                cell.fill = MANUAL_FILL
    ws.freeze_panes = "B2"


def export_register(mem: Memory, path: Path) -> dict:
    manual = _sync_manual(mem, _read_manual(path))
    plans = mem.latest_plans()
    pending = {}
    for a in mem.actions(("proposed", "approved", "failed")):
        if (a["status"] == "proposed" and (a["requires_approval"] or a["action_type"] == "ask_clarification")) \
                or (a["status"] != "proposed" and a["action_type"] == "draft_nudge"):
            pending.setdefault(a["commitment_id"], []).append(f"{a['action_id']} {a['action_type']} ({a['status']})")

    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    ws.append(REGISTER_COLUMNS)
    commitments = mem.all_commitments()
    for c in commitments:
        versions = mem.versions(c.commitment_id)
        superseded = [f"{v['due_date'] or '?'} (v{v['version_no']}, {v['approval']})" for v in versions
                      if not v["is_current"]]
        history = mem.conn.execute("SELECT * FROM actions WHERE commitment_id = ? ORDER BY created_at",
                                   (c.commitment_id,)).fetchall()
        nudges = sum(1 for a in history if a["action_type"] == "draft_nudge" and a["status"] in ("sent", "unsent_draft"))
        last_contact = next((f"{a['action_type']} {a['status']} {a['sent_at'] or a['created_at']}"
                             for a in reversed(history) if a["action_type"] == "draft_nudge"), "")
        plan = plans.get(c.commitment_id)
        ws.append([
            c.commitment_id, "I owe" if c.direction == "i_owe" else "Owed to me", c.owner, c.counterparty,
            c.description, c.due_date or "unresolved", c.due_text, "; ".join(superseded), c.status,
            c.status_reason or "", plan["action"] if plan else "", plan["decided_by"] if plan else "",
            plan["reason"] if plan else "", f"{c.source_message_id}: \"{c.evidence_quote}\"", nudges, last_contact,
            "; ".join(pending.get(c.commitment_id, [])), *manual.get(c.commitment_id, ("", "")), utc_now(),
        ])
    manual_idx = tuple(REGISTER_COLUMNS.index(m) + 1 for m in MANUAL_COLUMNS)
    _style(ws, {1: 18, 2: 11, 3: 24, 4: 24, 5: 40, 6: 13, 7: 18, 8: 26, 9: 15, 10: 30, 11: 16, 12: 11,
                13: 45, 14: 50, 15: 8, 16: 30, 17: 30, 18: 18, 19: 30, 20: 20}, manual_idx)

    hist = wb.create_sheet("Deadline history")
    hist.append(["Commitment ID", "Version", "Description", "Due date", "Due phrase", "Approval", "Current?",
                 "Change reason", "Source message", "Evidence quote", "Created (UTC)", "Superseded (UTC)"])
    for c in commitments:
        for v in mem.versions(c.commitment_id):
            hist.append([c.commitment_id, v["version_no"], v["description"], v["due_date"], v["due_text"],
                         v["approval"], "yes" if v["is_current"] else "no", v["change_reason"],
                         v["source_message_id"], v["evidence_quote"], v["created_at"], v["superseded_at"]])
    _style(hist, {1: 18, 2: 8, 3: 40, 4: 12, 5: 18, 6: 14, 7: 8, 8: 35, 9: 40, 10: 50, 11: 20, 12: 20})

    acts = wb.create_sheet("Actions")
    acts.append(["Action ID", "Cycle", "Commitment", "Type", "Status", "Needs approval", "Recipient",
                 "Delivered to", "Attempts", "Last error", "Rationale", "Created (UTC)", "Decided", "Sent"])
    for a in mem.actions():
        acts.append([a["action_id"], a["cycle_id"], a["commitment_id"], a["action_type"], a["status"],
                     "yes" if a["requires_approval"] else "no", a["recipient"], a["delivered_to"], a["attempts"],
                     a["last_error"], a["rationale"], a["created_at"], a["decided_at"], a["sent_at"]])
    _style(acts, {1: 16, 2: 14, 3: 16, 4: 22, 5: 13, 6: 9, 7: 24, 8: 24, 9: 8, 10: 30, 11: 50, 12: 20,
                  13: 20, 14: 20})

    about = wb.create_sheet("About")
    for line in [
        ["CommitKeeper review register", f"profile: {mem.profile_id}", f"exported {utc_now()}"],
        ["SQLite (data/commitkeeper.db) is the authoritative store; this workbook is a regenerated view."],
        ["Only the yellow columns 'Your decision' and 'Your comments' are read back on the next export."],
        ["Agent suggestions and your decisions are stored separately and never overwrite each other."],
    ]:
        about.append(line)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"path": str(path), "rows": len(commitments), "manual_entries": sum(1 for v in manual.values() if any(v))}


def export_ics(mem: Memory, path: Path) -> dict:
    """All-day events for active commitments. Stable UID + SEQUENCE = version, so calendar
    apps update an event when a deadline moves instead of creating a duplicate."""
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//CommitKeeper//EN", "CALSCALE:GREGORIAN"]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    count = 0
    for c in mem.active_commitments():
        if not c.due_date:
            continue
        due = date.fromisoformat(c.due_date)
        who = "You owe" if c.direction == "i_owe" else f"{c.counterparty} owes you"
        summary = f"{who}: {c.description}".replace(",", "\\,").replace(";", "\\;")
        lines += ["BEGIN:VEVENT", f"UID:{c.commitment_id}@commitkeeper", f"SEQUENCE:{c.version_no}",
                  f"DTSTAMP:{stamp}", f"DTSTART;VALUE=DATE:{due:%Y%m%d}",
                  f"DTEND;VALUE=DATE:{due + timedelta(days=1):%Y%m%d}", f"SUMMARY:{summary[:150]}",
                  f"DESCRIPTION:Evidence {c.source_message_id} (status {c.status})", "END:VEVENT"]
        count += 1
    lines.append("END:VCALENDAR")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")
    return {"path": str(path), "events": count}
